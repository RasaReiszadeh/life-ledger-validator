import csv
import datetime as dt
import os
import sqlite3
from openpyxl import load_workbook

FILE = next((f for f in ("Life_Ledger_Validator.xlsm", "Life_Ledger_Validator.xlsx") if os.path.exists(f)), None)
if FILE is None:
    raise SystemExit("Put Life_Ledger_Validator.xlsm (or .xlsx) in this folder first.")
print("Reading", FILE)
wb = load_workbook(FILE, data_only=True)

con = sqlite3.connect("ledger.db")
cur = con.cursor()
cur.executescript("""
    DROP TABLE IF EXISTS transactions;
    DROP TABLE IF EXISTS budget;
    CREATE TABLE transactions (
        row_num INTEGER, tx_date TEXT, description TEXT,
        category TEXT, type TEXT, amount REAL, raw_amount TEXT
    );
    CREATE TABLE budget (category TEXT, monthly_limit REAL);
""")

for i, row in enumerate(wb["Transactions"].iter_rows(min_row=2, max_col=5, values_only=True), start=2):
    d, desc, cat, typ, amt = row
    if all(v is None for v in row):
        continue
    iso = d.strftime("%Y-%m-%d") if isinstance(d, (dt.datetime, dt.date)) else None
    num = float(amt) if isinstance(amt, (int, float)) and not isinstance(amt, bool) else None
    cur.execute("INSERT INTO transactions VALUES (?,?,?,?,?,?,?)",
                (i, iso, desc, cat, typ, num, None if amt is None else str(amt)))

for cat, lim in wb["Budget"].iter_rows(min_row=2, max_col=2, values_only=True):
    if cat is not None:
        cur.execute("INSERT INTO budget VALUES (?,?)", (cat, lim))
con.commit()

issues = []


def flag(check, row, detail):
    issues.append((check, row, detail))
    print(f"  [{check}] row {row}: {detail}")


print("Invalid or missing dates")
for r in cur.execute("SELECT row_num, description FROM transactions WHERE tx_date IS NULL"):
    flag("date", r[0], r[1])

print("Non-numeric or missing amounts")
for r in cur.execute("SELECT row_num, description, raw_amount FROM transactions WHERE amount IS NULL"):
    flag("amount", r[0], f"{r[1]} -> {r[2]}")

print("Expense categories not in Budget")
for r in cur.execute("""
    SELECT t.row_num, t.description, t.category FROM transactions t
    LEFT JOIN budget b ON LOWER(t.category) = LOWER(b.category)
    WHERE t.type = 'Expense' AND b.category IS NULL
"""):
    flag("category", r[0], f"{r[1]} -> {r[2]}")

print("Possible duplicates")
for r in cur.execute("""
    SELECT tx_date, description, amount, COUNT(*), GROUP_CONCAT(row_num)
    FROM transactions WHERE tx_date IS NOT NULL AND amount IS NOT NULL
    GROUP BY tx_date, LOWER(description), amount HAVING COUNT(*) > 1
"""):
    flag("duplicate", r[4], f"{r[1]} on {r[0]} ({r[3]} times)")

print("Negative amounts")
for r in cur.execute("SELECT row_num, description, amount FROM transactions WHERE amount < 0"):
    flag("negative", r[0], f"{r[1]} -> {r[2]}")

print("\nSpend vs budget (SQL) compared with Summary sheet")
summary = wb["Summary"]
sheet_spent = {summary.cell(r, 1).value: summary.cell(r, 2).value for r in range(9, 16)}
for cat, lim, spent in cur.execute("""
    SELECT b.category, b.monthly_limit, COALESCE(SUM(t.amount), 0)
    FROM budget b LEFT JOIN transactions t
      ON LOWER(t.category) = LOWER(b.category) AND t.type = 'Expense'
    GROUP BY b.category ORDER BY b.rowid
"""):
    sheet = sheet_spent.get(cat)
    ok = sheet is not None and abs(spent - sheet) < 0.01
    print(f"  {cat:14} SQL {spent:9.2f} | sheet {sheet} | {'MATCH' if ok else 'MISMATCH'}")
    if not ok:
        flag("reconcile", cat, f"SQL {spent:.2f} vs sheet {sheet}")
    if spent > lim:
        flag("over_budget", cat, f"{spent - lim:.2f} over the limit of {lim:.2f}")

income = cur.execute("SELECT COALESCE(SUM(amount),0) FROM transactions WHERE type='Income'").fetchone()[0]
expense = cur.execute("SELECT COALESCE(SUM(amount),0) FROM transactions WHERE type='Expense'").fetchone()[0]
reported = summary["B5"].value
print(f"\nSQL income {income:.2f} | SQL expenses {expense:.2f} | SQL net {income - expense:.2f}")
if reported is None or abs((income - expense) - reported) > 0.01:
    flag("reconcile", "Summary!B5", f"reported net {reported} vs SQL net {income - expense:.2f}")

with open("reconciliation_report.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["check", "reference", "detail"])
    w.writerows(issues)
print(f"\n{len(issues)} findings written to reconciliation_report.csv")
